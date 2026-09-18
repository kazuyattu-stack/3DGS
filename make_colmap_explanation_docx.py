"""COLMAPとは何かを説明するWord文書(colmap_explanation.docx)を作成する。"""

from docx import Document
from docx.shared import Pt, Cm
from docx.oxml.ns import qn

doc = Document()

style = doc.styles["Normal"]
style.font.name = "Yu Gothic"
style.font.size = Pt(11)
rpr = style.element.get_or_add_rPr()
rFonts = rpr.find(qn("w:rFonts"))
if rFonts is None:
    rFonts = rpr.makeelement(qn("w:rFonts"), {})
    rpr.append(rFonts)
rFonts.set(qn("w:eastAsia"), "Yu Gothic")


def add_code_block(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9.5)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rPr.makeelement(qn("w:rFonts"), {})
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), "Consolas")
    return p


doc.add_heading("COLMAPとは", level=0)
doc.add_paragraph(
    "COLMAPは、複数枚の写真(または動画から抽出したフレーム)を入力として、"
    "「カメラがどこから・どの向きで撮影したか」と「被写体の3D形状」を同時に復元する、"
    "オープンソースの汎用SfM/MVSパイプラインソフトウェアです。"
    "スマホの写真や動画だけで3Dスキャンができる、いわゆる「フォトグラメトリ」「3D Gaussian Splatting」"
    "系のワークフローで、事実上の業界標準ツールとして広く使われています。"
)

doc.add_heading("開発元・ライセンス", level=1)
doc.add_paragraph(
    "スイス連邦工科大学チューリッヒ校(ETH Zurich)のJohannes Schönberger氏を中心に開発された、"
    "学術研究発のオープンソースソフトウェアです。BSDライセンスで公開されており、商用・非商用問わず"
    "無料で利用できます。"
)
doc.add_paragraph("・公式サイト: https://colmap.github.io/")
doc.add_paragraph("・GitHub: https://github.com/colmap/colmap")
doc.add_paragraph(
    "・元になった論文: \n"
    "  Schönberger and Frahm, “Structure-from-Motion Revisited”, CVPR 2016\n"
    "  Schönberger et al., “Pixelwise View Selection for Unstructured Multi-View Stereo”, ECCV 2016"
)

doc.add_heading("COLMAPが解く2つの問題", level=1)

doc.add_heading("① Structure-from-Motion(SfM)— カメラ姿勢とスパース点群の推定", level=2)
doc.add_paragraph(
    "複数の画像だけを手がかりに、「各カメラが3D空間のどこにあり、どちらを向いていたか」と"
    "「対応する特徴点が3D空間のどこにあったか」を同時に推定する処理です。次のステップで構成されます。"
)
for i, s in enumerate([
    "特徴点抽出: 各画像からSIFT特徴点を検出する",
    "特徴点マッチング: 画像同士で対応する特徴点を見つける",
    "Incremental SfM: 2枚の画像から復元を始め、1枚ずつ画像を追加しながら"
    "カメラ位置・向き・内部パラメータ(焦点距離など)・3D点群をバンドル調整で最適化していく",
], 1):
    doc.add_paragraph(f"{i}. {s}", style="List Number")
doc.add_paragraph(
    "出力は「カメラパラメータ」と「スパース(疎)な3D点群」です。点の数は数万点程度で、"
    "被写体の細かい形状までは分かりませんが、正確なカメラ位置・向きが得られるのが最大の価値です。"
)

doc.add_heading("② Multi-View Stereo(MVS)— 密な3D形状の復元", level=2)
doc.add_paragraph(
    "SfMで得たカメラパラメータを使い、画素レベルで密な奥行き(深度)を推定して、"
    "数十万〜数百万点規模の密な3D点群を作る処理です。GPU(CUDA)が必要です。"
)
for i, s in enumerate([
    "image_undistorter: レンズ歪みを補正した画像を作る",
    "patch_match_stereo: 各画像について、近傍視点との対応から画素ごとの深度マップを推定する",
    "stereo_fusion: 複数視点の深度マップを統合し、1つの密な3D点群にまとめる",
], 1):
    doc.add_paragraph(f"{i}. {s}", style="List Number")

doc.add_heading("主なコマンド", level=1)
table = doc.add_table(rows=1, cols=2)
table.style = "Light Grid Accent 1"
hdr = table.rows[0].cells
hdr[0].text = "コマンド"
hdr[1].text = "役割"
for cmd, desc in [
    ("automatic_reconstructor", "特徴点抽出〜SfM〜(必要なら)MVSまでを一括で自動実行する“おまかせ”コマンド"),
    ("feature_extractor", "画像からSIFT特徴点を抽出する"),
    ("sequential_matcher / exhaustive_matcher", "特徴点マッチング。動画のような連番画像には\nsequential(近傍フレームのみ照合)が高速"),
    ("mapper", "Incremental SfM本体。カメラ姿勢とスパース点群を推定する"),
    ("image_undistorter", "レンズ歪みを補正した画像とカメラパラメータを出力する"),
    ("patch_match_stereo", "GPU(CUDA)を使って密な深度マップを推定する"),
    ("stereo_fusion", "複数視点の深度マップを統合して密な点群(fused.ply)を作る"),
    ("model_converter", "出力形式をBIN/TXT/PLYの間で変換する"),
    ("gui", "GUIで結果を可視化・編集するビューア"),
]:
    cells = table.add_row().cells
    cells[0].text, cells[1].text = cmd, desc

doc.add_heading("入力と出力", level=1)
doc.add_paragraph("・入力: 画像ファイル群(JPEG/PNG等)。動画ファイルは直接読み込めないため、事前にフレーム抽出が必要")
doc.add_paragraph("・出力(SfM): カメラの内部/外部パラメータ、スパース点群(cameras.bin, images.bin, points3D.bin)")
doc.add_paragraph("・出力(MVS、任意): 密な点群(fused.ply)")

doc.add_heading("何に使われているか", level=1)
doc.add_paragraph(
    "COLMAPそのものは3Dスキャン・測量・文化財のデジタルアーカイブ・VR/ARコンテンツ制作など"
    "幅広い用途で使われていますが、近年特に注目されているのが、NeRFや3D Gaussian Splattingといった"
    "ニューラル3D表現の学習における「前処理」としての役割です。"
    "これらの手法は、写真だけから3Dシーンを学習しますが、その学習には「各写真がどのカメラ位置・"
    "向きで撮られたか」という正確な情報が必要になります。COLMAPのSfMはこの情報を提供する"
    "デファクトスタンダードのツールになっており、3D Gaussian Splatting公式実装も"
    "COLMAP形式のデータセットをそのまま読み込めるように作られています。"
)

doc.add_heading("今回のパイプラインでの役割", level=1)
doc.add_paragraph(
    "今回のワークフローでは、動画から抽出したフレーム群に対してCOLMAPを2段階で使っています。"
)
doc.add_paragraph("① SfM: カメラ姿勢とスパース点群を推定(3D Gaussian Splattingの学習に必須)", style="List Bullet")
doc.add_paragraph("② MVS: 密な点群を生成(任意。学習には不要だが、比較用に生成)", style="List Bullet")
doc.add_paragraph(
    "つまりCOLMAPは、動画という「見た目の情報」を、3D Gaussian Splattingが学習できる"
    "「カメラ位置つきの構造化データ」に変換する、パイプライン全体の要となる工程を担っています。"
)

doc.add_heading("参考リンク", level=1)
for link in [
    "公式サイト: https://colmap.github.io/",
    "GitHub: https://github.com/colmap/colmap",
    "ドキュメント: https://colmap.github.io/tutorial.html",
]:
    doc.add_paragraph(link, style="List Bullet")

doc.save("../colmap_explanation.docx")
print("saved colmap_explanation.docx")
